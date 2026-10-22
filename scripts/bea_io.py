"""BEA input-output: who an industry buys from, by SEC SIC code — writes pipeline/data/bea_io.json (pipeline/ioflow.py).

Run once a year when BEA publishes new tables (needs openpyxl; LibreOffice to read Census's old .xls concordances):

  python scripts/bea_io.py --bea AllTablesIO.zip --census DIR --fixed-assets detailnonres_inv1.xlsx [--out …]

  --bea     https://apps.bea.gov/industry/iTables%20Static%20Files/AllTablesIO.zip (the Use and Make tables)
  --census  a folder with Census's concordances, converted to .xlsx where they are .xls
            (soffice --headless --convert-to xlsx *.xls):
              https://www.census.gov/naics/concordances/1987_SIC_to_1997_NAICS.xls
              https://www.census.gov/naics/concordances/1997_NAICS_to_2002_NAICS.xls
              https://www.census.gov/naics/concordances/2002_to_2007_NAICS.xls
              https://www.census.gov/naics/concordances/2007_to_2012_NAICS.xls
              https://www.census.gov/naics/concordances/2012_to_2017_NAICS.xlsx

What it keeps, for each of BEA's ~400 detail industries (the 2017 benchmark Use table, before redefinitions, producer
prices — the finest table BEA publishes; the summary table is annual but has only 71 industries):
  - how one dollar of its operating spending divides between purchased inputs, employee compensation and taxes on
    production (value added's other part, gross operating surplus, is profit and depreciation — not spending)
  - its purchased inputs by supplying industry: each commodity bought goes to imports at that commodity's import share
    of total supply, the rest to the domestic industries that make it in proportion to their output of it (the Make
    table: the industry-technology assumption statistical agencies use)
  - the same for capital spending, by its own mix of assets (BEA fixed assets: investment by industry and asset type,
    https://apps.bea.gov/national/FA2004/Details/xls/detailnonres_inv1.xlsx — computers, trucks, office buildings,
    power structures …, each asset mapped to the commodities it is); the economy's investment mix where an industry
    has none (government)
  - its gross output, moved from 2017 to the latest year by its summary industry's growth (the denominator a listed
    supplier's revenue share is taken of)
and for every SEC SIC code the BEA industries it corresponds to (1987 SIC → 1997 NAICS → … → 2017 NAICS → BEA), weighted
by their output. Every figure downstream is therefore an estimate at industry-average precision.
"""
import argparse
import io
import json
import os
import re
import zipfile
from collections import defaultdict

import openpyxl

MIN_SHARE = 0.002            # suppliers below 0.2% of a buyer's inputs go to "other"

# BEA fixed assets: investment by industry (sheet) and asset type, detailnonres_inv1.xlsx — the industries' NAICS
FA_NAICS = {'110C': ['111', '112'], '113F': ['113', '114', '115'], '2110': ['211'], '2120': ['212'], '2130': ['213'],
            '2211': ['2211'], '2212': ['2212'], '2213': ['2213'], '2300': ['23'], '3210': ['321'], '3270': ['327'],
            '3310': ['331'], '3320': ['332'], '3330': ['333'], '3340': ['334'], '3350': ['335'],
            '336M': ['3361', '3362', '3363'], '336O': ['3364', '3365', '3366', '3369'], '3370': ['337'], '3380': ['3391'],
            '3390': ['3399'], '3110': ['311'], '3120': ['312'], '313T': ['313', '314'], '315A': ['315', '316'],
            '3220': ['322'], '3230': ['323'], '3240': ['324'], '3250': ['325'], '3260': ['326'], '4210': ['423'],
            '4220': ['424', '425'], '4410': ['441'], '4450': ['445'], '4520': ['452'],
            '442R': ['442', '443', '444', '446', '447', '448', '451', '453', '454', '4A'], '4810': ['481'], '4820': ['482'],
            '4830': ['483'], '4840': ['484'], '4850': ['485'], '4860': ['486'], '487S': ['487', '488', '492'],
            '4930': ['493'], '5110': ['511', '516'], '5120': ['512'], '5130': ['515', '517'], '5140': ['518', '519'],
            '5210': ['521'], '5221': ['5221'], '5223': ['5223'], '5229': ['5222'], '5230': ['523'], '524A': ['5241'],
            '5243': ['5242'], '5250': ['525'], '5310': ['531'], '5320': ['532', '533'], '5411': ['5411'],
            '5415': ['5415'], '5412': ['5412', '5413', '5414', '5416', '5417', '5418', '5419'], '5500': ['55'],
            '5610': ['561'], '5620': ['562'], '6100': ['61'], '6210': ['621'], '622H': ['622'], '6230': ['623'],
            '6240': ['624'], '711A': ['711', '712'], '7130': ['713'], '7210': ['721'], '7220': ['722'], '8100': ['81']}
# what a company's capital spending (property, plant and equipment, purchased software) buys: BEA's asset types → the
# detail commodities that are those assets. Own-account software, R&D and entertainment originals are left out: a
# company's capex line does not include them (they are its own staff's work, expensed).
ASSETS = {'EP1A': {'334111': 1}, 'EP1B': {'334111': 1}, 'EP1C': {'334112': 1}, 'EP1D': {'334118': 1},
          'EP1E': {'334118': 1}, 'EP1F': {'334112': 1}, 'EP1G': {'334112': 1}, 'EP1H': {'541512': 1},
          'EP20': {'334210': .3, '334220': .5, '334290': .2}, 'EP34': {'339112': .6, '339113': .3, '339114': .1},
          'EP35': {'334510': .8, '334517': .2}, 'EP36': {'334513': .3, '334515': .2, '334516': .2, '334511': .2, '334512': .1},
          'EP31': {'333316': 1}, 'EP12': {'333318': .5, '339940': .5}, 'EI11': {'325180': 1},
          'EI12': {'332310': .4, '332420': .2, '332999': .4}, 'EI21': {'333611': 1}, 'EI22': {'333618': 1},
          'EI30': {'333517': .5, '333514': .25, '33351B': .25}, 'EI40': {'33329A': .7, '333242': .3},
          'EI50': {'333912': .15, '333914': .15, '333920': .25, '333415': .25, '33399A': .2},
          'EI60': {'335311': .5, '335313': .3, '335314': .2}, 'ET11': {'336112': 1},
          'ET12': {'336120': .7, '336212': .2, '336211': .1}, 'ET20': {'336111': 1}, 'ET30': {'336411': .7, '336412': .3},
          'ET40': {'336611': .7, '336612': .3}, 'ET50': {'336500': 1}, 'EO11': {'337121': .5, '337122': .5},
          'EO12': {'33721A': .7, '337215': .3}, 'EO30': {'333111': 1}, 'EO21': {'333111': 1}, 'EO40': {'333120': 1},
          'EO22': {'333120': 1}, 'EO50': {'333130': 1}, 'EO60': {'333318': 1}, 'EO71': {'335220': 1},
          'EO72': {'335999': 1}, 'EO80': {'339990': .5, '332999': .25, '33399A': .25},
          'SOO1': {'2332A0': 1}, 'SB31': {'233210': 1}, 'SB32': {'233210': 1}, 'SOO2': {'233210': 1}, 'SC03': {'2332A0': 1},
          'SC04': {'2332A0': 1}, 'SC01': {'2332A0': 1}, 'SOMO': {'3219A0': 1}, 'SC02': {'2332A0': 1}, 'SI00': {'233230': 1},
          'SU30': {'233240': 1}, 'SU60': {'233240': 1}, 'SU40': {'233240': 1}, 'SU50': {'233240': 1}, 'SU20': {'233240': 1},
          'SM01': {'213111': 1}, 'SM02': {'2332D0': 1}, 'SB10': {'2332D0': 1}, 'SB20': {'233262': 1}, 'SB41': {'2332A0': 1},
          'SB42': {'2332D0': 1}, 'SB43': {'2332C0': 1}, 'SB45': {'2332C0': 1}, 'SU11': {'2332C0': 1}, 'SU12': {'2332C0': 1},
          'SB44': {'2332C0': 1}, 'SB46': {'2332C0': 1}, 'SN00': {'2332D0': 1}, 'SO01': {'2332D0': 1}, 'SO02': {'2332D0': 1},
          'SO03': {'2332D0': 1}, 'SO04': {'2332C0': 1}, 'ENS1': {'511200': 1}, 'ENS2': {'541511': 1}}


# SIC codes where the concordance chain lands on the wrong BEA industry (old NAICS codes split many ways, minor pieces in
# large industries): the largest SEC codes by revenue, checked by hand. A list sums to one.
SIC_OVERRIDE = {
    '7370': {'519130': .45, '518200': .25, '511200': .15, '541512': .15},     # Alphabet, Meta: search, ads, cloud
    '7389': {'541610': .3, '561900': .3, '5419A0': .2, '541512': .2},       # Accenture, Uber, Visa: services n.e.c.
    '6324': {'5241XX': 1}, '6331': {'5241XX': 1},                             # health, property & casualty insurers
    '3576': {'334210': 1},                                                    # Cisco, Arista: data communications gear
    '6199': {'522A00': 1},                                                    # American Express: card lending
    '4700': {'561500': 1},                                                    # Booking, Expedia: travel arrangement
    '4400': {'483000': 1},                                                    # cruise lines, shipping
    '3600': {'335312': .25, '333611': .25, '335311': .2, '335313': .15, '335999': .15},
    '3570': {'334111': .5, '334118': .2, '334112': .1, '541512': .2},        # IBM, HP: computers and services
    '3841': {'339112': .7, '339113': .3},
    '3531': {'333120': .8, '333130': .1, '333920': .1},                       # Caterpillar
    '1531': {'233411': .7, '233412': .15, '2334A0': .15},                     # home builders
    '5411': {'445000': 1}, '5500': {'441000': .85, '447000': .15},
    '2070': {'311224': .6, '311225': .4},
}


def division(sic):
    """The NAICS sectors a SIC division corresponds to."""
    d = int(sic[:2])
    for lo, hi, naics in ((1, 9, ('11',)), (10, 14, ('21',)), (15, 17, ('23',)), (20, 39, ('31', '32', '33')),
                          (40, 47, ('48', '49')), (48, 48, ('51', '48')), (49, 49, ('22', '562', '486')), (50, 51, ('42',)),
                          (52, 57, ('44', '45')), (58, 58, ('722',)), (59, 59, ('44', '45')),
                          (60, 64, ('52',)), (65, 65, ('53',)), (67, 67, ('52', '55', '53')),
                          (70, 89, ('51', '53', '54', '56', '61', '62', '71', '72', '81'))):
        if lo <= d <= hi:
            return naics
    return ()


def fixed_assets(path, years=3):
    """{industry: ({asset: $M, average of the last `years` years}, title, last year)} from detailnonres_inv1.xlsx."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    out = {}
    for name in wb.sheetnames:
        if name not in FA_NAICS:
            continue
        rows = [list(r) for r in wb[name].iter_rows(values_only=True)]
        head = next(r for r in rows if r and str(r[0] or '').strip() == 'Asset Codes')
        cols = [j for j, c in enumerate(head) if str(c or '').strip().isdigit()][-years:]
        assets = {}
        for r in rows:
            code = str(r[0] or '').strip()
            if code in ASSETS:
                vals = [num(r[j]) for j in cols if j < len(r)]
                assets[code] = sum(vals) / len(vals) if vals else 0.0
        out[name] = (assets, str(rows[0][0] or name), str(head[cols[-1]]))
    return out


def rows_of(data, sheet=None):
    wb = openpyxl.load_workbook(io.BytesIO(data) if isinstance(data, bytes) else data, read_only=True, data_only=True)
    ws = wb[sheet] if sheet else wb.worksheets[0]
    return [list(r) for r in ws.iter_rows(values_only=True)]


def num(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return 0.0


def table(rows, code_row=5):
    """(column codes, {row code: {column code: value}}, names) from a BEA Use or Make sheet."""
    head = rows[code_row]
    names = {}
    cols = {}
    for j, c in enumerate(head):
        if c and j >= 2:
            cols[j] = str(c).strip()
    for j, n in enumerate(rows[code_row - 1]):
        if j in cols and n:
            names[cols[j]] = str(n).strip()
    out = {}
    for r in rows[code_row + 1:]:
        if not r or not r[0]:
            continue
        code = str(r[0]).strip()
        names.setdefault(code, str(r[1] or '').strip())
        out[code] = {cols[j]: num(r[j]) for j in cols if j < len(r) and num(r[j])}
    return list(cols.values()), out, names


def bea_tables(zpath):
    z = zipfile.ZipFile(zpath)
    pick = lambda pat: next(n for n in z.namelist() if re.search(pat, n))
    use = rows_of(z.read(pick(r'^IOUse_Before_Redefinitions_PRO_2017_Detail')), '2017')
    make = rows_of(z.read(pick(r'^IOMake_Before_Redefinitions_2017_Detail')), '2017')
    naics = rows_of(z.read(pick(r'^IOUse_Before_Redefinitions_PRO_2017_Detail')), 'NAICS Codes')
    summ_name = pick(r'^IOUse_Before_Redefinitions_PRO_1997-\d{4}_Summary')
    swb = openpyxl.load_workbook(io.BytesIO(z.read(summ_name)), read_only=True, data_only=True)
    years = sorted(int(s) for s in swb.sheetnames if s.isdigit())
    summary = {}
    for y in (2017, years[-1]):
        _, t, _ = table([list(r) for r in swb[str(y)].iter_rows(values_only=True)])
        summary[y] = t.get('T008') or {}                     # gross output by summary industry
    return use, make, naics, summary, years[-1]


def detail_naics(rows):
    """{detail code: (summary code, [2017 NAICS prefixes], title)} from the Use workbook's 'NAICS Codes' sheet."""
    out, summary = {}, None
    for r in rows[5:]:
        r = (list(r) + [None] * 8)[:8]
        if r[1] and str(r[1]).strip():
            summary = str(r[1]).strip()
        code = str(r[3] or '').strip()
        if not code:
            continue
        rel = re.findall(r'\d{2,6}', str(r[6] or '').replace('-', ' '))
        span = re.findall(r'(\d{2,6})\s*-\s*(\d{2,6})', str(r[6] or ''))
        for a, b in span:                                   # "4231-4239" style ranges
            if len(a) == len(b):
                rel += [str(x) for x in range(int(a), int(b) + 1)]
        out[code] = (summary, sorted(set(rel)), str(r[4] or '').strip())
    return out


def concordance(path, a_col, b_col, skip):
    """{from code: set(to codes)} from one Census concordance sheet."""
    m = defaultdict(set)
    for r in rows_of(path)[skip:]:
        if len(r) <= max(a_col, b_col):
            continue
        a, b = str(r[a_col] or '').strip(), str(r[b_col] or '').strip()
        if re.fullmatch(r'\d{2,6}', a) and re.fullmatch(r'\d{2,6}', b):
            m[a].add(b)
    return m


def chain_naics(census):
    """SIC 1987 (4-digit) → {2017 NAICS code: weight}."""
    sic = defaultdict(set)
    for r in rows_of(os.path.join(census, '1987_SIC_to_1997_NAICS.xlsx'))[1:]:
        s, n = str(r[0] or '').strip(), str(r[3] or '').strip()
        if re.fullmatch(r'\d{4}', s) and re.fullmatch(r'\d{5,6}', n):
            sic[s].add(n)
    steps = [concordance(os.path.join(census, '1997_NAICS_to_2002_NAICS.xlsx'), 1, 0, 2),
             concordance(os.path.join(census, '2002_to_2007_NAICS.xlsx'), 0, 2, 3),
             concordance(os.path.join(census, '2007_to_2012_NAICS.xlsx'), 0, 2, 3),
             concordance(os.path.join(census, '2012_to_2017_NAICS.xlsx'), 0, 2, 3)]

    def step(codes, m):
        out = defaultdict(float)
        for c, w in codes.items():
            hit = m.get(c) or (set().union(*(v for k, v in m.items() if k.startswith(c))) if len(c) < 6 else None)
            if not hit:
                out[c] += w                                  # unchanged by this revision
                continue
            for h in hit:
                out[h] += w / len(hit)
        return out
    res = {}
    for s, ns in sic.items():
        codes = {n: 1 / len(ns) for n in ns}
        for m in steps:
            codes = step(codes, m)
        res[s] = dict(codes)
    return res


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--bea', required=True)
    ap.add_argument('--census', required=True)
    ap.add_argument('--fixed-assets', required=True, help='https://apps.bea.gov/national/FA2004/Details/xls/detailnonres_inv1.xlsx')
    ap.add_argument('--out', default=os.path.join(os.path.dirname(__file__), '..', 'pipeline', 'data', 'bea_io.json'))
    a = ap.parse_args()
    use_rows, make_rows, naics_rows, summary, last_year = bea_tables(a.bea)
    ucols, U, names = table(use_rows)
    _, M, mnames = table(make_rows)
    names.update({k: v for k, v in mnames.items() if k not in names})
    det = detail_naics(naics_rows)
    industries = [c for c in ucols if not re.match(r'^(T0|F0|F1)', c)]
    commodities = [c for c in U if not re.match(r'^(T0|V0)', c)]
    # commodity c → producing industries (Make: industry rows × commodity columns), and its import share
    made = defaultdict(dict)
    for ind, row in M.items():
        if re.match(r'^(T0|V0)', ind):
            continue
        for c, v in row.items():
            if v > 0 and not c.startswith('T0'):
                made[c][ind] = v
    imp_share = {}
    for c in commodities:
        dom = sum(made[c].values())
        imp = -U[c].get('F05000', 0.0)
        imp_share[c] = max(0.0, min(1.0, imp / (dom + imp))) if dom + imp > 0 else (1.0 if c == 'S00300' else 0.0)

    def to_suppliers(purchases):
        """{commodity: $} → ({supplier industry: share}, import share): imports first, then the producers' shares."""
        tot = sum(v for v in purchases.values() if v > 0)
        if tot <= 0:
            return {}, 0.0
        sup, imports = defaultdict(float), 0.0
        for c, v in purchases.items():
            if v <= 0:
                continue
            m = imp_share.get(c, 0.0)
            imports += v * m
            prod = made.get(c) or {}
            ptot = sum(prod.values())
            if ptot <= 0:
                sup['_OTHER'] += v * (1 - m)
                continue
            for i, q in prod.items():
                sup[i] += v * (1 - m) * q / ptot
        sup = {k: v / tot for k, v in sup.items()}
        small = sum(v for k, v in sup.items() if v < MIN_SHARE or k == '_OTHER')
        sup = {k: round(v, 5) for k, v in sup.items() if v >= MIN_SHARE and k != '_OTHER'}
        if small > 0:
            sup['_OTHER'] = round(small, 5)
        return sup, round(imports / tot, 5)

    # output growth 2017 → latest by summary industry
    grow = {}
    for s, v in summary[2017].items():
        w = summary[last_year].get(s)
        if v and w:
            grow[s] = w / v
    out_ind = {}
    for b in industries:
        col = {c: U[c].get(b, 0.0) for c in commodities}
        inter = sum(v for v in col.values() if v > 0)
        comp = U.get('V00100', {}).get(b, 0.0)
        tax = max(0.0, U.get('V00200', {}).get(b, 0.0))
        gross = U.get('T008', {}).get(b, 0.0)
        spend = inter + comp + tax
        if gross <= 0 or spend <= 0:
            continue
        sup, imp = to_suppliers(col)
        summ = (det.get(b) or (None,))[0]
        out_ind[b] = {'name': names.get(b, b), 'summary': summ,
                      'output': round(gross * 1e6 * grow.get(summ, 1.0)),       # dollars, latest year
                      'inter': round(inter / spend, 5), 'comp': round(comp / spend, 5), 'tax': round(tax / spend, 5),
                      'spend_out': round(spend / gross, 5),                                # spending per dollar of output
                      'imports': imp, 'suppliers': sup}
    inv = defaultdict(float)
    for c in commodities:
        for f in ('F02E00', 'F02N00', 'F02S00'):
            inv[c] += max(0.0, U[c].get(f, 0.0))
    inv_sup, inv_imp = to_suppliers(inv)
    # capital spending by industry: its own asset mix (BEA fixed assets) → commodities → suppliers and imports
    capex, fa_year = {}, None
    for fa, (assets, title, year) in fixed_assets(a.fixed_assets).items():
        buy = defaultdict(float)
        for code, v in assets.items():
            for c, w in ASSETS[code].items():
                if c in U:
                    buy[c] += v * w
        sup, imp = to_suppliers(buy)
        if sup:
            capex[fa] = {'name': re.sub(r'^\S+\s*-\s*', '', title).strip(), 'suppliers': sup, 'imports': imp}
            fa_year = year
    missing = sorted({c for m in ASSETS.values() for c in m if c not in U})
    if missing:
        print('asset commodities not in the Use table:', missing)

    def fa_of(b):
        rel = (det.get(b) or (None, [b[:6]]))[1] or [re.sub(r'\D.*$', '', b)]
        best, size = None, 0
        for fa, pre in FA_NAICS.items():
            for p in pre:
                if any(r.startswith(p) for r in rel) and len(p) > size:
                    best, size = fa, len(p)
        return best if best in capex else None
    for b in out_ind:
        fa = fa_of(b)
        if fa:
            out_ind[b]['fa'] = fa
    # SIC → BEA detail industries
    sic_naics = chain_naics(a.census)
    bea_naics = {b: rel for b, (_, rel, _) in det.items() if b in out_ind}

    def bea_for(n):
        hits = [b for b, rel in bea_naics.items() if any(n.startswith(r) or r.startswith(n) for r in rel)]
        if not hits:
            return {}
        # the most specific match wins (a 6-digit NAICS beats its 3-digit parent's catch-all)
        best = max(max(len(r) for r in bea_naics[b] if n.startswith(r) or r.startswith(n)) for b in hits)
        hits = [b for b in hits if any((n.startswith(r) or r.startswith(n)) and len(r) == best for r in bea_naics[b])]
        tot = sum(out_ind[b]['output'] for b in hits) or 1
        return {b: out_ind[b]['output'] / tot for b in hits}
    sic_map = {}
    for s, ns in sic_naics.items():
        # every BEA industry one of the SIC code's NAICS codes falls in, weighted by its output: the concordance lists
        # each NAICS piece of a SIC code once without its size (prepackaged software is mostly software publishing,
        # hardly at all reproducing disks)
        # only the pieces in the SIC code's own part of the economy: a minor piece in another sector can hit a large BEA
        # industry and outweigh the rest (ship building's ship-repair piece is NAICS 488390, "support activities for
        # transportation", and would make General Dynamics a sightseeing operator)
        own = [n for n in ns if any(n.startswith(p) for p in division(s))] or list(ns)
        acc = {}
        for n in own:
            for b in bea_for(n):
                acc[b] = out_ind[b]['output']
        tot = sum(acc.values())
        if tot > 0:
            sic_map[s] = [[b, round(v / tot, 4)] for b, v in sorted(acc.items(), key=lambda kv: -kv[1]) if v / tot >= 0.01]
    for s, m in SIC_OVERRIDE.items():
        sic_map[s] = [[b, w] for b, w in m.items() if b in out_ind]
    doc = {'source': {'use': 'BEA Use table before redefinitions, producer prices, 2017 detail (402 industries)',
                      'make': 'BEA Make table before redefinitions, 2017 detail',
                      'scale': f'BEA summary Use tables, gross output 2017 → {last_year}',
                      'concordance': 'Census 1987 SIC → 1997 NAICS → 2002 → 2007 → 2012 → 2017 NAICS; BEA detail ↔ 2017 NAICS',
                      'url': 'https://www.bea.gov/industry/input-output-accounts-data'},
           'year': 2017, 'scale_year': last_year, 'industries': out_ind,
           'invest': {'suppliers': inv_sup, 'imports': inv_imp,
                      'note': 'private nonresidential fixed investment (equipment, structures, intellectual property), 2017'},
           'capex': capex, 'capex_note': f'BEA fixed assets: investment by industry and asset type, average of {int(fa_year) - 2}–{fa_year}; '
                                         'equipment, structures and purchased software (own-account software, R&D and '
                                         'entertainment originals left out: not in a company\'s capex line)',
           'sic': sic_map}
    os.makedirs(os.path.dirname(os.path.abspath(a.out)), exist_ok=True)
    with open(a.out, 'w', encoding='utf-8') as f:
        json.dump(doc, f, separators=(',', ':'), sort_keys=True)
    print(f'{len(out_ind)} BEA industries, {len(sic_map)} SIC codes mapped, scale year {last_year} → {a.out} '
          f'({os.path.getsize(a.out) / 1e3:.0f} KB)')


if __name__ == '__main__':
    main()
