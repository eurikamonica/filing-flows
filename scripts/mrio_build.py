"""Inter-country input-output coefficients for Money flow Beta 1.1 → pipeline/data/mrio.json.

Three free multi-country tables, reduced to what the beta needs for each country and industry: how its output divides
into purchased inputs, employee pay, taxes and value added, and where its purchased inputs come from (by country and by
supplying industry). Run by .github/workflows/mrio-refresh.yml (the sources are too large to keep in the repository):

  FIGARO   Eurostat's EU inter-country tables, industry by industry (naio_10_fcp_ii4): the EU, EFTA, candidate
           countries and the main trading partners (United States, China, Japan, India, Brazil …), 64 industries,
           the latest year; the only one of the three that splits value added (employee pay, other taxes, surplus)
  ADB      the Asian Development Bank's multiregional tables (kidb.adb.org): 74 economies and the rest of the world,
           35 sectors, the latest year
  ICIO     the OECD's inter-country tables (oe.cd/icio): about 80 economies, 50 industries. Its download page sits
           behind a browser check, so the zip is used only when it has been put on the repository's release "icio-data"
           by hand (see the workflow); without it the file is built from the other two

  python scripts/mrio_build.py --figaro-year 2024 --adb adb.xlsx [--icio icio.zip --icio-year 2022] --out mrio.json
"""
import argparse
import csv
import gzip
import io
import json
import os
import re
import sys
import time
import urllib.request
import zipfile
from collections import defaultdict

TOP_ORIGINS, TOP_INDS = 14, 18

# ISIC Rev. 4 sections → divisions
SECTIONS = {"A": range(1, 4), "B": range(5, 10), "C": range(10, 34), "D": [35], "E": range(36, 40), "F": range(41, 44),
            "G": range(45, 48), "H": range(49, 54), "I": range(55, 57), "J": range(58, 64), "K": range(64, 67), "L": [68],
            "M": range(69, 76), "N": range(77, 83), "O": [84], "P": [85], "Q": range(86, 89), "R": range(90, 94),
            "S": range(94, 97), "T": range(97, 99), "U": [99]}
# ADB MRIO's 35 sectors (the WIOD classification, ISIC Rev. 3 read as Rev. 4 divisions)
ADB_SECTORS = {"c1": [1, 2, 3], "c2": [5, 6, 7, 8, 9], "c3": [10, 11, 12], "c4": [13, 14], "c5": [15], "c6": [16],
               "c7": [17, 18], "c8": [19], "c9": [20, 21], "c10": [22], "c11": [23], "c12": [24, 25], "c13": [28],
               "c14": [26, 27], "c15": [29, 30], "c16": [31, 32, 33], "c17": [35, 36], "c18": [41, 42, 43], "c19": [45],
               "c20": [46], "c21": [47], "c22": [55, 56], "c23": [49], "c24": [50], "c25": [51], "c26": [52, 79],
               "c27": [53, 61], "c28": [64, 65, 66], "c29": [68], "c30": [62, 63, 69, 70, 71, 72, 73, 74, 75, 77, 78, 80, 81, 82],
               "c31": [84], "c32": [85], "c33": [86, 87, 88], "c34": [37, 38, 39, 58, 59, 60, 90, 91, 92, 93, 94, 95, 96],
               "c35": [97, 98]}
ADB_ISO = {"SWI": "CHE", "PRC": "CHN", "GER": "DEU", "DEN": "DNK", "SPA": "ESP", "UKG": "GBR", "INO": "IDN", "IRE": "IRL",
           "NET": "NLD", "POR": "PRT", "ROM": "ROU", "TAP": "TWN", "BAN": "BGD", "MAL": "MYS", "PHI": "PHL", "VIE": "VNM",
           "MON": "MNG", "SRI": "LKA", "FIJ": "FJI", "BRU": "BRN", "BHU": "BTN", "CAM": "KHM", "MLD": "MDV", "NEP": "NPL",
           "SIN": "SGP", "KUW": "KWT", "UAE": "ARE", "TAJ": "TJK", "KGZ": "KGZ", "RoW": "ROW", "ROW": "ROW"}
ISO2 = {"AL": "ALB", "AR": "ARG", "AT": "AUT", "AU": "AUS", "BE": "BEL", "BG": "BGR", "BR": "BRA", "CA": "CAN", "CH": "CHE",
        "CN": "CHN", "CY": "CYP", "CZ": "CZE", "DE": "DEU", "DK": "DNK", "EE": "EST", "EL": "GRC", "ES": "ESP", "FI": "FIN",
        "FR": "FRA", "HR": "HRV", "HU": "HUN", "ID": "IDN", "IE": "IRL", "IN": "IND", "IT": "ITA", "JP": "JPN", "KR": "KOR",
        "LT": "LTU", "LU": "LUX", "LV": "LVA", "ME": "MNE", "MK": "MKD", "MT": "MLT", "MX": "MEX", "NL": "NLD", "NO": "NOR",
        "PL": "POL", "PT": "PRT", "RO": "ROU", "RS": "SRB", "RU": "RUS", "SA": "SAU", "SE": "SWE", "SI": "SVN", "SK": "SVK",
        "TR": "TUR", "UK": "GBR", "US": "USA", "ZA": "ZAF", "WRL_REST": "ROW"}
FIG_FINAL = {"P3_S13", "P3_S14", "P3_S15", "P51G", "P5M"}
FIG_VA = {"D1": "comp", "D29X39": "otax", "B2A3G": "gos", "D21X31": "ptax"}


def divisions(code):
    """ISIC divisions of a table's industry code: C26 → [26], C10-12 / C10T12 → 10..12, C31_32 → 31, 32, J → 58..63,
    C24A → [24], C302T309 → [30]."""
    c = code.strip().upper()
    m = re.match(r"^([A-U])(.*)$", c)
    if not m:
        return []
    sec, rest = m.group(1), m.group(2)
    if not rest:
        return list(SECTIONS.get(sec, []))
    out = []
    for part in rest.split("_"):
        r = re.match(r"^(\d{2})\d*[A-Z]?(?:[-T](\d{2})\d*)?$", part)
        if r:
            a = int(r.group(1))
            b = int(r.group(2)) if r.group(2) else a
            out += list(range(a, b + 1))
    return out


def get(url, path=None, ua=None, tries=4):
    for k in range(tries):
        try:
            req = urllib.request.Request(url, headers={"User-Agent": ua or "filing-flows mrio build"})
            with urllib.request.urlopen(req, timeout=600) as r:
                data = r.read()
            if path:
                open(path, "wb").write(data)
            return data
        except Exception as e:                                  # noqa: BLE001 — retried, then raised
            if k == tries - 1:
                raise
            print(f"  retry {url}: {e}", file=sys.stderr)
            time.sleep(10 * (k + 1))


class Acc:
    """Column sums of one table: for each buying (country, industry), its inputs by origin and by supplying industry,
    and the value-added rows."""

    def __init__(self):
        self.org = defaultdict(lambda: defaultdict(float))
        self.ind = defaultdict(lambda: defaultdict(float))
        self.inter = defaultdict(float)
        self.va = defaultdict(lambda: defaultdict(float))
        self.out = defaultdict(float)

    def add_inter(self, buyer, origin, ind, v):
        if v:
            self.org[buyer][origin] += v
            self.ind[buyer][ind] += v
            self.inter[buyer] += v

    def compact(self, codes):
        res = {}
        for b, o in self.out.items():
            if o <= 0:
                continue
            iso, code = b
            inter = self.inter.get(b, 0.0)
            org, ind = self.org.get(b, {}), self.ind.get(b, {})
            rec = {"o": round(o, 1), "i": round(inter / o, 5)}
            for k, v in self.va.get(b, {}).items():
                rec[k] = round(v / o, 5)
            if inter > 0:
                rec["dom"] = round(org.get(iso, 0.0) / inter, 5)
                top = sorted(org.items(), key=lambda kv: -kv[1])
                rec["org"] = {k: round(v / inter, 5) for k, v in top[:TOP_ORIGINS] if v > 0}
                topi = sorted(ind.items(), key=lambda kv: -kv[1])
                rec["ind"] = {k: round(v / inter, 5) for k, v in topi[:TOP_INDS] if v > 0}
            res[f"{iso}|{code}"] = rec
        return res


def figaro(year, countries=None):
    """Eurostat naio_10_fcp_ii4, one destination country per request (the API refuses more than 5 million rows)."""
    base = "https://ec.europa.eu/eurostat/api/dissemination/sdmx/2.1/data/naio_10_fcp_ii4"
    acc, codes = Acc(), {}
    dests = countries or sorted(k for k in ISO2 if k != "WRL_REST") + ["WRL_REST"]
    for d in dests:
        url = f"{base}/A...{d}..?format=SDMX-CSV&startPeriod={year}&endPeriod={year}"
        try:
            raw = get(url)
        except Exception as e:                                  # noqa: BLE001
            print(f"::warning::FIGARO {d}: {e}")
            continue
        n = 0
        for row in csv.DictReader(io.StringIO(raw.decode("utf-8"))):
            use, ava, orig = row["ind_use"], row["ind_ava"], row["c_orig"]
            if use in FIG_FINAL or not row["OBS_VALUE"]:
                continue
            v = float(row["OBS_VALUE"])
            b = (ISO2.get(d, d), use)
            codes.setdefault(use, divisions(use))
            if ava in FIG_VA:
                acc.va[b][FIG_VA[ava]] += v
                acc.out[b] += v
            elif ava.startswith("OP_"):
                continue
            else:
                acc.add_inter(b, ISO2.get(orig, orig), ava, v)
                acc.out[b] += v
            n += 1
        print(f"  FIGARO {d}: {n} cells")
    return {"year": year, "currency": "EUR", "codes": codes, "t": acc.compact(codes)}


def adb(path):
    """ADB MRIO workbook: rows = supplying (country, sector), columns = buying (country, sector) then final demand."""
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True)
    ws = wb.worksheets[0]
    title = " ".join(str(r[0]) for r in ws.iter_rows(min_row=1, max_row=3, values_only=True) if r[0])
    rows = ws.iter_rows(values_only=True)
    head = [next(rows) for _ in range(7)]
    ctry, sect = head[5], head[6]
    cols = {}
    for j in range(4, len(sect)):
        s, c = sect[j], ctry[j]
        if isinstance(s, str) and re.fullmatch(r"c\d{1,2}", s) and c:
            cols[j] = (ADB_ISO.get(c, c), s)
    acc = Acc()
    year = (re.search(r"(20\d\d)", title) or [None, None])[1]
    for r in rows:
        code, orig = r[3], r[2]
        if not code:
            continue
        code = str(code)
        if re.fullmatch(r"c\d{1,2}", code):
            o = ADB_ISO.get(orig, orig)
            for j, b in cols.items():
                v = r[j]
                if isinstance(v, (int, float)) and v:
                    acc.add_inter(b, o, code, float(v))
        elif code in ("r99", "r64", "r69"):
            key = {"r99": "ptax", "r64": "va", "r69": None}[code]
            for j, b in cols.items():
                v = r[j]
                if isinstance(v, (int, float)):
                    if key:
                        acc.va[b][key] += float(v)
                    else:
                        acc.out[b] = float(v)
    return {"year": int(year) if year else None, "currency": "USD", "title": title.strip(),
            "codes": {k: v for k, v in ADB_SECTORS.items()}, "t": acc.compact(ADB_SECTORS)}


def icio(path, year):
    """OECD ICIO csv (one year) inside the zip: rows and columns '<ISO3>_<industry>', then TLS, VA and OUT rows."""
    z = zipfile.ZipFile(path)
    name = next(n for n in z.namelist() if str(year) in n and n.lower().endswith(".csv"))
    f = io.TextIOWrapper(z.open(name), encoding="utf-8")
    r = csv.reader(f)
    head = next(r)
    cols = {}
    for j, h in enumerate(head[1:], 1):
        m = re.match(r"^([A-Z0-9]{3})_(.+)$", h)
        if m and divisions(m.group(2)):
            iso = {"MX1": "MEX", "MX2": "MEX", "CN1": "CHN", "CN2": "CHN"}.get(m.group(1), m.group(1))
            cols[j] = (iso, m.group(2))
    acc, codes = Acc(), {}
    for row in r:
        lab = row[0]
        m = re.match(r"^([A-Z0-9]{3})_(.+)$", lab)
        if m and divisions(m.group(2)):
            o = {"MX1": "MEX", "MX2": "MEX", "CN1": "CHN", "CN2": "CHN"}.get(m.group(1), m.group(1))
            codes.setdefault(m.group(2), divisions(m.group(2)))
            for j, b in cols.items():
                v = row[j]
                if v and v != "0":
                    acc.add_inter(b, o, m.group(2), float(v))
        elif lab in ("TLS", "VA", "OUT") or lab.endswith(("_TLS", "_VA")):
            key = "ptax" if "TLS" in lab else "va" if "VA" in lab else None
            for j, b in cols.items():
                v = float(row[j] or 0)
                if key:
                    acc.va[b][key] += v
                else:
                    acc.out[b] += v
    return {"year": year, "currency": "USD", "codes": codes, "t": acc.compact(codes)}


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--figaro-year", type=int, default=2024)
    ap.add_argument("--figaro-countries", help="comma-separated ISO2 destinations (default: all)")
    ap.add_argument("--adb", help="ADB MRIO workbook (xlsx)")
    ap.add_argument("--icio", help="OECD ICIO zip (optional)")
    ap.add_argument("--icio-year", type=int, default=2022)
    ap.add_argument("--edgar", help="EDGAR state/country codes page (html), for foreign filers' home countries")
    ap.add_argument("--out", default=os.path.join(os.path.dirname(__file__), "..", "pipeline", "data", "mrio.json"))
    a = ap.parse_args()
    doc = {"generated": time.strftime("%Y-%m-%d"), "sources": {}}
    t0 = time.time()
    doc["sources"]["figaro"] = figaro(a.figaro_year, a.figaro_countries.split(",") if a.figaro_countries else None)
    print(f"FIGARO: {len(doc['sources']['figaro']['t'])} country-industries ({time.time() - t0:.0f} s)")
    if a.adb and os.path.exists(a.adb):
        doc["sources"]["adb"] = adb(a.adb)
        print(f"ADB: {len(doc['sources']['adb']['t'])} country-industries ({time.time() - t0:.0f} s)")
    if a.icio and os.path.exists(a.icio):
        doc["sources"]["icio"] = icio(a.icio, a.icio_year)
        print(f"ICIO: {len(doc['sources']['icio']['t'])} country-industries ({time.time() - t0:.0f} s)")
    if a.edgar and os.path.exists(a.edgar):
        t = open(a.edgar, encoding="utf-8", errors="replace").read()
        doc["edgar"] = {c: n.strip() for c, n in re.findall(r"<td[^>]*>\s*([A-Z0-9]{2})\s*</td>\s*<td[^>]*>\s*([^<]+?)\s*</td>", t)}
    with open(a.out, "w", encoding="utf-8") as f:
        json.dump(doc, f, separators=(",", ":"))
    print(f"wrote {a.out}: {os.path.getsize(a.out) / 1e6:.1f} MB")


if __name__ == "__main__":
    main()
